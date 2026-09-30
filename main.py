from openpyxl.styles import Border
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.styles import Side
import requests
# pyrefly: ignore [missing-import]
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import PatternFill
resposta = requests.get('https://www.cemig.com.br/valores-e-tarifas/tarifas-vigentes/')

if(resposta.status_code == 200):
    soup = BeautifulSoup(resposta.content, 'html.parser')
 
    tabelas = soup.find_all('table')
    
    tabelas_tarifa_convencional = []
    for tabela in tabelas:
        tabelas_tarifa_convencional.append(tabela)
    
    honeydewfill = PatternFill(start_color="F0FFF0", end_color="F0FFF0", fill_type="solid")
    whitesmokefill = PatternFill(start_color="DCDCDC", end_color="DCDCDC", fill_type="solid")

    wb = Workbook()
    del wb['Sheet']
    sheet = wb.create_sheet("Tabelas de tarifas")
    linha = 1
    linha_index = 1
    coluna = 1

    for tabela in tabelas:
        for linha_atual in tabela.find_all('tr'):
            for coluna_atual in linha_atual.find_all(['td', 'th']):
                if linha_index == linha:
                    sheet.cell(linha,coluna).fill = honeydewfill
                if(coluna == 1 and linha != linha_index):
                    sheet.cell(linha, coluna).fill = whitesmokefill

                sheet.cell(linha, coluna).value = coluna_atual.get_text(strip=True)
                sheet.cell(linha, coluna).alignment = Alignment(horizontal='left', vertical='center')
                sheet.cell(linha, coluna).border = Border(left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'), top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
                letra_coluna = get_column_letter(coluna)
                sheet.column_dimensions[letra_coluna].width = 40
                
                coluna += 1
            linha += 1
            coluna = 1
        linha += 2
        linha_index = linha

    wb.save("teste.xlsx")



    
    

